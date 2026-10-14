"""Génération Word (.docx) de l'acte de procuration ministérielle."""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from apps.procurations.pdf import _load_procuration, _template_context

if TYPE_CHECKING:
    from apps.procurations.models import Procuration

INK = RGBColor(0x11, 0x11, 0x11)
SLATE = RGBColor(0x33, 0x33, 0x33)
WARN = RGBColor(0x8B, 0x1A, 0x1A)
FONT = "Calibri"
BODY_SIZE = Pt(14)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_DISCLAIMER = (
    "Copie de travail AvoLex — lecture seule. "
    "Les données ci-dessous sont extraites de la base ; "
    "toute modification de ce fichier n'est PAS enregistrée dans AvoLex. "
    "L'acte officiel reste le PDF généré dans l'application."
)


def build_procuration_docx(procuration: Procuration) -> bytes:
    """Rend l'acte de procuration en Word (copie de travail, lecture seule)."""
    procuration = _load_procuration(procuration)
    ctx = _template_context(procuration)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.4)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = BODY_SIZE
    style.font.color.rgb = INK
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = notice.add_run(_DISCLAIMER)
    run.italic = True
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.color.rgb = WARN
    notice.paragraph_format.space_after = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("// PROCURATION //")
    run.bold = True
    run.font.size = BODY_SIZE
    run.font.name = FONT
    run.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(4)

    number = doc.add_paragraph()
    number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = number.add_run(f"N° {ctx['document_number']}")
    run.font.size = BODY_SIZE
    run.font.name = FONT
    number.paragraph_format.space_after = Pt(6)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    pPr = rule._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "double")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "111111")
    pBdr.append(bottom)
    pPr.append(pBdr)

    preamble = doc.add_paragraph()
    preamble.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    preamble.paragraph_format.first_line_indent = Cm(1.0)
    _add_runs(
        preamble,
        [
            ("La ", False),
            ("REPUBLIQUE DEMOCRATIQUE DU CONGO", True),
            (", représentée par le ", False),
            (str(ctx["minister_title"]), False),
            (", ", False),
            (str(ctx["minister_name"]), True),
            (", nommé par ", False),
            (str(ctx["ordinance_appointment"]), False),
            (", ayant son Cabinet sis à ", False),
            (str(ctx["ministry_address"]), False),
            (", agissant en vertu de ", False),
            (str(ctx["ordinance_attributions"]), False),
            (",", False),
        ],
    )

    mandate = doc.add_paragraph()
    mandate.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    mandate.paragraph_format.first_line_indent = Cm(1.0)
    mandate_parts: list[tuple[str, bool]] = [
        ("donne par les présentes ", False),
        ("mandat à", True),
        (" ", False),
    ]
    if ctx["lawyer_name"]:
        mandate_parts.append((str(ctx["lawyer_name"]), True))
        if ctx["lawyer_bar"]:
            mandate_parts.append((f", {ctx['lawyer_bar']}", False))
        if ctx["lawyer_address"]:
            mandate_parts.append((f", dont le Cabinet est établi {ctx['lawyer_address']}", False))
        mandate_parts.append((",", False))
    elif ctx["cabinet"]:
        cab = ctx["cabinet"]
        mandate_parts.append(("le cabinet ", False))
        mandate_parts.append((cab.name, True))  # type: ignore[union-attr]
        if cab.bar_association:  # type: ignore[union-attr]
            mandate_parts.append((f", {cab.bar_association}", False))
        if ctx["cabinet_address"]:
            mandate_parts.append(
                (f", dont le Cabinet est établi {ctx['cabinet_address']}", False)
            )
        mandate_parts.append((",", False))
    else:
        mandate_parts.append(("l'avocat désigné,", False))
    mandate_parts.append(
        (" agissant conjointement ou séparément avec pouvoir de substitution.", False)
    )
    _add_runs(mandate, mandate_parts)

    _labeled(doc, "Aux fins de :", str(ctx["purpose_text"]))
    _labeled(doc, "En cause :", str(ctx["en_cause"]))
    _labeled(doc, "Contre :", str(ctx["contre"]))

    fait = doc.add_paragraph()
    fait.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fait.paragraph_format.space_before = Pt(8)
    run = fait.add_run(f"Fait à Kinshasa, le {ctx['granted_date']}")
    run.font.name = FONT
    run.font.size = BODY_SIZE

    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign.paragraph_format.space_before = Pt(28)
    run = sign.add_run(str(ctx["signatory_capacity"]))
    run.font.name = FONT
    run.font.size = BODY_SIZE

    if ctx["signatory_name"]:
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        name_p.paragraph_format.space_before = Pt(28)
        run = name_p.add_run(str(ctx["signatory_name"]))
        run.bold = True
        run.font.name = FONT
        run.font.size = BODY_SIZE

    if ctx["signatory_title"]:
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = title_p.add_run(str(ctx["signatory_title"]))
        run.font.name = FONT
        run.font.size = BODY_SIZE

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(36)
    lines = []
    if ctx["granted_by"]:
        lines.append(str(ctx["granted_by"]))
    lines.append(f"Le {ctx['printed_at']}")
    lines.append("Visa Dircaba")
    run = footer.add_run("\n".join(lines))
    run.font.name = FONT
    run.font.size = Pt(11)
    run.font.color.rgb = SLATE

    _protect_read_only(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _protect_read_only(doc: Document) -> None:
    """Empêche les modifications accidentelles (désactivable volontairement dans Word)."""
    protection = OxmlElement("w:documentProtection")
    protection.set(qn("w:edit"), "readOnly")
    protection.set(qn("w:enforcement"), "1")
    doc.settings.element.append(protection)


def _add_runs(paragraph: object, parts: list[tuple[str, bool]]) -> None:
    for text, bold in parts:
        run = paragraph.add_run(text)  # type: ignore[attr-defined]
        run.bold = bold
        run.font.name = FONT
        run.font.size = BODY_SIZE
        run.font.color.rgb = INK


def _labeled(doc: Document, label: str, body: str) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.space_before = Pt(6)
    run = para.add_run(f"{label} ")
    run.bold = True
    run.font.name = FONT
    run.font.size = BODY_SIZE
    run.font.color.rgb = INK
    run = para.add_run(body)
    run.font.name = FONT
    run.font.size = BODY_SIZE
    run.font.color.rgb = INK


def procuration_docx_filename(procuration: Procuration) -> str:
    """Nom de fichier Word sûr pour le téléchargement."""
    ref = (procuration.reference or "acte").replace("/", "-").replace("\\", "-")
    return f"procuration-{ref}.docx"
